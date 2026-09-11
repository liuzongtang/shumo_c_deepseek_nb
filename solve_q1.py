# -*- coding: utf-8 -*-
"""问题1：理想单日计划购电策略（纯线性规划）。
每天电价/负载相同，光伏用预测数据，储能 0:00 与 24:00 电量相同。
"""
import numpy as np
from scipy.optimize import linprog
from data_loader import load_fj1, DT, PMAX_E, E_MAX, E_MIN, E_INIT, ETA, N_SLOT


def solve_day(price, load, pv, E0=E_INIT, E1=E_INIT):
    """解一天 144 段的储能+购电 LP。
    price: (144,) 元/kWh；load/pv: (144,) kW。
    返回 dict: G,C,D,E (E 长度 145), 购电费(元)。"""
    Le = load * DT      # kWh
    PVe = pv * DT       # kWh

    # 变量顺序: G[0..143], C[0..143], D[0..143], E[0..144]
    nG, nC, nD, nE = N_SLOT, N_SLOT, N_SLOT, N_SLOT + 1
    offG, offC, offD, offE = 0, nG, nG + nC, nG + nC + nD
    n = offE + nE

    c = np.zeros(n)
    c[offG:offG + nG] = price            # 目标: min Σ G*price

    Aeq = []
    beq = []
    Aub = []
    bub = []
    # (1) 功率平衡: G + PV·Δt + D >= L·Δt + C  =>  -G - D + C <= (PV-L)·Δt
    for t in range(N_SLOT):
        row = np.zeros(n)
        row[offG + t] = -1.0
        row[offD + t] = -1.0
        row[offC + t] = 1.0
        Aub.append(row); bub.append(PVe[t] - Le[t])
    # (2) 储能动态: E[t+1] - E[t] - η*C[t] + D[t]/η = 0
    for t in range(N_SLOT):
        row = np.zeros(n)
        row[offE + t + 1] = 1.0
        row[offE + t] = -1.0
        row[offC + t] = -ETA
        row[offD + t] = 1.0 / ETA
        Aeq.append(row); beq.append(0.0)
    # (3) 周期/初末: E[0]=E0, E[144]=E1
    row = np.zeros(n); row[offE + 0] = 1.0; Aeq.append(row); beq.append(E0)
    row = np.zeros(n); row[offE + N_SLOT] = 1.0; Aeq.append(row); beq.append(E1)

    Aeq = np.array(Aeq); beq = np.array(beq)
    Aub = np.array(Aub); bub = np.array(bub)

    # 边界
    lb = np.zeros(n); ub = np.full(n, np.inf)
    for t in range(N_SLOT):
        ub[offC + t] = PMAX_E          # 充电 ≤ 833.33
        ub[offD + t] = PMAX_E          # 放电 ≤ 833.33
    for t in range(nE):
        lb[offE + t] = E_MIN
        ub[offE + t] = E_MAX
    # G 无上界 (lb=0, ub=inf)

    res = linprog(c, A_eq=Aeq, b_eq=beq, A_ub=Aub, b_ub=bub,
                  bounds=list(zip(lb, ub)), method="highs")
    if not res.success:
        raise RuntimeError(res.message)

    G = res.x[offG:offG + nG]
    C = res.x[offC:offC + nC]
    D = res.x[offD:offD + nD]
    E = res.x[offE:offE + nE]
    cost = float(np.dot(G, price))
    return {'G': G, 'C': C, 'D': D, 'E': E, 'cost': cost}


def main():
    d = load_fj1()
    price, load, pv = d['price'], d['load'], d['pv']
    r = solve_day(price, load, pv)

    print("电价范围: %.4f ~ %.4f 元/kWh" % (price.min(), price.max()))
    print("负载范围: %.0f ~ %.0f kW" % (load.min(), load.max()))
    print("光伏峰值: %.0f kW" % pv.max())
    print("全天购电量: %.2f kWh" % r['G'].sum())
    print("全天购电费: %.2f 元" % r['cost'])
    print("0:00 储电量: %.2f kWh   24:00 储电量: %.2f kWh" % (r['E'][0], r['E'][-1]))
    print("储能最低/最高: %.2f / %.2f kWh" % (r['E'].min(), r['E'].max()))
    print("总充电: %.2f kWh  总放电: %.2f kWh" % (r['C'].sum(), r['D'].sum()))

    # ---- 表1：指定时段购电量 ----
    print("\n表1  微网在指定时间段的购电量")
    slots = {'10:00-10:10': 60, '12:00-12:10': 72, '14:00-14:10': 84,
             '16:00-16:10': 96, '18:00-18:10': 108, '20:00-20:10': 120}
    for k, t in slots.items():
        print("  %s: %.4f kWh" % (k, r['G'][t]))

    # ---- 表2：4小时时段充放电 + 0:00/24:00储电量 ----
    print("\n表2  储能设备充放电量及储电量")
    periods = [('0:00-4:00', 0, 24), ('4:00-8:00', 24, 48), ('8:00-12:00', 48, 72),
               ('12:00-16:00', 72, 96), ('16:00-20:00', 96, 120), ('20:00-24:00', 120, 144)]
    for name, a, b in periods:
        print("  %s: 充电 %.4f kWh, 放电 %.4f kWh" % (name, r['C'][a:b].sum(), r['D'][a:b].sum()))
    print("  0:00 储电量 %.4f kWh" % r['E'][0])
    print("  24:00 储电量 %.4f kWh" % r['E'][-1])
    return r


def write_result1(r, out_path="result1.xlsx"):
    """按附件5模板写 result1.xlsx。
    时间槽按标签对齐：模板列'0:10-0:20'填入时段[0:10,0:20]的值(循环平移一格)。
    """
    import openpyxl
    tpl = "附件/附件5/result1.xlsx"
    wb = openpyxl.load_workbook(tpl)

    # --- 计划购电量 sheet ---
    ws = wb["计划购电量"]
    n = 144
    for row in range(2, 2 + n):          # 数据行 2..145
        r0 = row - 2                    # 模板行下标 0..143
        val = r['G'][(r0 + 1) % n]       # 循环平移：'0:10-0:20' <- 时段[0:10,0:20]=slot1
        ws.cell(row=row, column=2, value=round(float(val), 4))

    # --- 充放电量 sheet ---
    ws = wb["充放电量"]
    periods = [(0, 24), (24, 48), (48, 72), (72, 96), (96, 120), (120, 144)]
    for i, (a, b) in enumerate(periods):
        row = 2 + i
        ws.cell(row=row, column=2, value=round(float(r['C'][a:b].sum()), 4))  # 充电量
        ws.cell(row=row, column=3, value=round(float(r['D'][a:b].sum()), 4))  # 放电量
    # 0:00 / 24:00 储电量 (在第1、2行的 时刻/储电量 列)
    ws.cell(row=2, column=4, value="0:00")
    ws.cell(row=2, column=5, value=round(float(r['E'][0]), 4))
    ws.cell(row=3, column=4, value="24:00")
    ws.cell(row=3, column=5, value=round(float(r['E'][-1]), 4))

    wb.save(out_path)
    print("已写出", out_path)


if __name__ == "__main__":
    r = main()
    write_result1(r)
