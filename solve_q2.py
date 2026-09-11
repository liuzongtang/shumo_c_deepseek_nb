# -*- coding: utf-8 -*-
"""问题2：全年计划购电策略（确定性 + 鲁棒两套方案）。

模型假设：
  - 每天电价相同(用附件1分时电价)，负载/光伏随时间变化(附件2实际数据)。
  - 储能设备跨日连续运行，但每天 0:00 制定当天计划、24:00 回到名义电量 6000 kWh
    （因电价逐日周期重复，跨日搬移无套利空间，此假设与全年最优一致）。
  - 平衡约束用不等式(允许光伏过剩时弃光)：G + PV·Δt + D >= L·Δt + C。

确定性方案：0:00 已知当天实际负载/光伏(附件2)，紧急购电=0。
鲁棒方案：  0:00 用附件3的0:00光伏预报(负载视为已知、逐日规律性强)制定计划 G_plan；
           实际揭晓后储能重调度最小化缺口，剩余缺口按 5 倍交易电价紧急购电。
"""
import numpy as np
import datetime as dt
from scipy.optimize import linprog
from data_loader import (load_fj1, load_fj2, load_fj3, DT, PMAX_E, E_MAX, E_MIN,
                         E_INIT, ETA, N_SLOT)

PRICE = None          # 附件1 电价 (144,)
DATES = None          # 附件2 日期字符串
LOAD = None           # 附件2 实际负载 (365,144)
PV = None             # 附件2 实际光伏 (365,144)


def solve_day(price, load, pv, E0=E_INIT, E1=E_INIT, E_min=E_MIN, E_max=E_MAX):
    """解一天 144 段的储能+购电 LP（平衡用不等式，允许弃光）。
    price:(144,) 元/kWh；load/pv:(144,) kW。返回 dict: G,C,D,E(145),cost。
    E_min/E_max 为储能 SOC 上下限(用于敏感性测试，默认用全局 E_MIN/E_MAX)。"""
    Le = load * DT
    PVe = pv * DT
    nG = nC = nD = N_SLOT
    nE = N_SLOT + 1
    offG, offC, offD, offE = 0, nG, nG + nC, nG + nC + nD
    n = offE + nE

    c = np.zeros(n)
    c[offG:offG + nG] = price

    # 等式约束：储能动态 + 初末电量
    Aeq, beq = [], []
    for t in range(N_SLOT):
        row = np.zeros(n)
        row[offE + t + 1] = 1.0
        row[offE + t] = -1.0
        row[offC + t] = -ETA
        row[offD + t] = 1.0 / ETA
        Aeq.append(row); beq.append(0.0)
    row = np.zeros(n); row[offE + 0] = 1.0; Aeq.append(row); beq.append(E0)
    row = np.zeros(n); row[offE + N_SLOT] = 1.0; Aeq.append(row); beq.append(E1)

    # 不等式约束：G + D - C >= (L-PV)·Δt  =>  -G - D + C <= (PV-L)·Δt
    Aub, bub = [], []
    for t in range(N_SLOT):
        row = np.zeros(n)
        row[offG + t] = -1.0
        row[offD + t] = -1.0
        row[offC + t] = 1.0
        Aub.append(row); bub.append(PVe[t] - Le[t])

    lb = np.zeros(n); ub = np.full(n, np.inf)
    for t in range(N_SLOT):
        ub[offC + t] = PMAX_E
        ub[offD + t] = PMAX_E
    for t in range(nE):
        lb[offE + t] = E_min
        ub[offE + t] = E_max

    res = linprog(c, A_eq=np.array(Aeq), b_eq=np.array(beq),
                  A_ub=np.array(Aub), b_ub=np.array(bub),
                  bounds=list(zip(lb, ub)), method="highs")
    if not res.success:
        raise RuntimeError(res.message)
    return {
        'G': res.x[offG:offG + nG],
        'C': res.x[offC:offC + nC],
        'D': res.x[offD:offD + nD],
        'E': res.x[offE:offE + nE],
        'cost': float(np.dot(res.x[offG:offG + nG], price)),
    }


def build_pv_forecast():
    """从附件3 0:00预报构造 (365,144) 的10分钟光伏预报(step-hold，整点值平铺6段)。"""
    fc = load_fj3()
    base = dt.datetime(2025, 1, 1)
    pv_f = np.zeros((365, N_SLOT))
    for i in range(365):
        d = base + dt.timedelta(days=i)
        key = "%d-%d-%d" % (d.year, d.month, d.day)     # '2025-1-1' 格式
        hourly = fc.get((key, '0:00'))
        if hourly is None:
            raise KeyError(key)
        for h in range(24):
            pv_f[i, h * 6:(h + 1) * 6] = hourly[h]
    return pv_f


def recourse(G_plan, load, pv, E0=E_INIT, E1=E_INIT):
    """给定固定计划购电量 G_plan，对实际负载/光伏重调度储能，最小化紧急购电。
    返回 C,D,E(重调度后), e(紧急购电量 kWh/段)。"""
    Le = load * DT
    PVe = pv * DT
    nC = nD = N_SLOT
    nE = N_SLOT + 1
    ne = N_SLOT
    offC, offD, offE = 0, nC, nC + nD
    offe = offE + nE
    n = offe + ne

    c = np.zeros(n)
    c[offe:offe + ne] = PRICE          # min Σ e·price (5倍为常数因子)

    Aeq, beq = [], []
    for t in range(N_SLOT):            # 储能动态
        row = np.zeros(n)
        row[offE + t + 1] = 1.0
        row[offE + t] = -1.0
        row[offC + t] = -ETA
        row[offD + t] = 1.0 / ETA
        Aeq.append(row); beq.append(0.0)
    row = np.zeros(n); row[offE + 0] = 1.0; Aeq.append(row); beq.append(E0)
    row = np.zeros(n); row[offE + N_SLOT] = 1.0; Aeq.append(row); beq.append(E1)

    Aub, bub = [], []
    for t in range(N_SLOT):            # G_plan+pv·Δt+D-C+e >= load·Δt
        row = np.zeros(n)
        row[offD + t] = -1.0
        row[offC + t] = 1.0
        row[offe + t] = -1.0
        Aub.append(row); bub.append(G_plan[t] + PVe[t] - Le[t])

    lb = np.zeros(n); ub = np.full(n, np.inf)
    for t in range(N_SLOT):
        # 充电量上限：储能只能充"实际盈余"(计划购电+实际光伏-负载)，不能从紧急购电套利
        surplus = G_plan[t] + PVe[t] - Le[t]
        ub[offC + t] = min(PMAX_E, max(0.0, surplus))
        ub[offD + t] = PMAX_E
    for t in range(nE):
        lb[offE + t] = E_MIN; ub[offE + t] = E_MAX

    res = linprog(c, A_eq=np.array(Aeq), b_eq=np.array(beq),
                  A_ub=np.array(Aub), b_ub=np.array(bub),
                  bounds=list(zip(lb, ub)), method="highs")
    if not res.success:
        raise RuntimeError(res.message)
    return {'C': res.x[offC:offC + nC], 'D': res.x[offD:offD + nD],
            'E': res.x[offE:offE + nE], 'e': res.x[offe:offe + ne]}


def run_deterministic(start=31, ndays=334):
    """确定性：每天用附件2实际数据求解，紧急购电=0。返回各天 G/C/D/E。"""
    Gs, Cs, Ds, Es = [], [], [], []
    for d in range(start, start + ndays):
        r = solve_day(PRICE, LOAD[d], PV[d])
        Gs.append(r['G']); Cs.append(r['C']); Ds.append(r['D']); Es.append(r['E'])
    return {'G': np.array(Gs), 'C': np.array(Cs), 'D': np.array(Ds), 'E': np.array(Es)}


def run_robust(start=31, ndays=334):
    """鲁棒：0:00 用附件3光伏预报(负载视为已知)制定计划 G_plan；
    实际揭晓后储能重调度、最小化紧急购电 e。返回 G_plan/C/D/E + 紧急购电 e。"""
    pv_f = build_pv_forecast()
    Gs, Cs, Ds, Es, Em = [], [], [], [], []
    for d in range(start, start + ndays):
        r = solve_day(PRICE, LOAD[d], pv_f[d])   # 计划：预报光伏 + 实际负载
        G_plan = r['G']
        rc = recourse(G_plan, LOAD[d], PV[d])    # 执行：对实际光伏重调度储能
        Gs.append(G_plan)
        Cs.append(rc['C']); Ds.append(rc['D']); Es.append(rc['E']); Em.append(rc['e'])
    return {'G': np.array(Gs), 'C': np.array(Cs), 'D': np.array(Ds),
            'E': np.array(Es), 'e': np.array(Em)}


def _slot_clock(t):
    """第 t 段(0-based)起始时刻 HH:MM。"""
    m = t * 10
    return "%02d:%02d" % (m // 60, m % 60)


def _merge_ranges(e):
    """把一天 144 段的紧急购电量合并成连续时间段列表 [(start_slot,end_slot,kWh)]。"""
    out = []
    i = 0
    while i < N_SLOT:
        if e[i] <= 1e-9:
            i += 1
            continue
        j = i
        while j + 1 < N_SLOT and e[j + 1] > 1e-9:
            j += 1
        out.append((i, j, float(e[i:j + 1].sum())))
        i = j + 1
    return out


def write_result2(res, out_path):
    """按附件5模板写 result2.xlsx。res 含 G(ndays,144), C, D, E, e(可为None)。
    计划购电量列：列(2+i) <- 数据 slot (i+1) mod 144 (与result1一致循环平移)。"""
    import openpyxl
    tpl = "附件/附件5/result2.xlsx"
    wb = openpyxl.load_workbook(tpl)
    ndays = res['G'].shape[0]
    G, C, D, E = res['G'], res['C'], res['D'], res['E']
    e = res.get('e')

    # ---- 计划购电量 sheet（模板已有 334 个日期行，直接填）----
    ws = wb["计划购电量"]
    for k in range(ndays):
        row = 2 + k
        g = G[k]
        for i in range(N_SLOT):
            ws.cell(row=row, column=2 + i,
                    value=round(float(g[(i + 1) % N_SLOT]), 4))
        ws.cell(row=row, column=2 + N_SLOT, value=round(float(g.sum()), 4))        # 全天购电量
        ws.cell(row=row, column=3 + N_SLOT, value=round(float(np.dot(g, PRICE)), 4))  # 全天购电费

    # ---- 充放电量 sheet（重建：334 天 × 6 时段）----
    ws = wb["充放电量"]
    ws.delete_rows(2, ws.max_row - 1)
    periods = [(0, 24, '0:00-4:00'), (24, 48, '4:00-8:00'), (48, 72, '8:00-12:00'),
               (72, 96, '12:00-16:00'), (96, 120, '16:00-20:00'), (120, 144, '20:00-24:00')]
    base = dt.datetime(2025, 2, 1)
    r = 2
    for k in range(ndays):
        for pi, (a, b, name) in enumerate(periods):
            if pi == 0:
                c1 = ws.cell(row=r, column=1, value=base + dt.timedelta(days=k))
                c1.number_format = 'mm-dd-yy'
            ws.cell(row=r, column=2, value=name)
            ws.cell(row=r, column=3, value=round(float(C[k][a:b].sum()), 4))
            ws.cell(row=r, column=4, value=round(float(D[k][a:b].sum()), 4))
            if pi == 0:
                c5 = ws.cell(row=r, column=5, value=dt.time(0, 0))
                c5.number_format = 'h:mm'
                ws.cell(row=r, column=6, value=round(float(E[k][0]), 4))
            elif pi == 1:
                c5 = ws.cell(row=r, column=5, value='24:00')
                c5.number_format = '@'
                ws.cell(row=r, column=6, value=round(float(E[k][-1]), 4))
            r += 1

    # ---- 紧急购电量 sheet（重建：仅 e>0 的天）----
    ws = wb["紧急购电量"]
    ws.delete_rows(2, ws.max_row - 1)
    r = 2
    for k in range(ndays):
        if e is None:
            continue
        ranges = _merge_ranges(e[k])
        if not ranges:
            continue
        for ri, (a, b, amt) in enumerate(ranges):
            if ri == 0:
                c1 = ws.cell(row=r, column=1, value=base + dt.timedelta(days=k))
                c1.number_format = 'mm-dd-yy'
            ws.cell(row=r, column=2, value="%s-%s" % (_slot_clock(a), _slot_clock(b + 1)))
            ws.cell(row=r, column=3, value=round(amt, 4))
            r += 1

    wb.save(out_path)
    print("已写出", out_path)


def _find_idx(date_str):
    """在附件2日期序列里找 'YYYY.M.D' 的 0-based 下标。"""
    for i, s in enumerate(DATES):
        if str(s).replace('-', '.') == date_str or str(s) == date_str:
            return i
    # 兜底：按日期解析
    target = dt.datetime.strptime(date_str, "%Y.%m.%d").date()
    for i, s in enumerate(DATES):
        if hasattr(s, 'date') and s.date() == target:
            return i
        if isinstance(s, (int, float)):
            continue
        try:
            if dt.datetime.strptime(str(s), "%Y-%m-%d").date() == target:
                return i
        except ValueError:
            pass
    raise KeyError(date_str)


def summary(res, label, e_mode="emergency"):
    G, C, D = res['G'], res['C'], res['D']
    e = res.get('e')
    plan_cost = float(np.sum(G * PRICE[None, :]))
    plan_kwh = float(G.sum())
    lines = ["===== %s =====" % label]
    lines.append("全年(2.1-12.31)计划购电量: %.2f kWh" % plan_kwh)
    lines.append("全年(2.1-12.31)计划购电费: %.2f 元" % plan_cost)
    if e is not None:
        em_kwh = float(e.sum())
        em_cost = float(np.sum(e * PRICE[None, :])) * 5.0
        lines.append("全年紧急购电量: %.2f kWh" % em_kwh)
        lines.append("全年紧急购电费(5倍): %.2f 元" % em_cost)
        lines.append("全年总购电费: %.2f 元" % (plan_cost + em_cost))
        n_em_days = int((e.sum(axis=1) > 1e-9).sum())
        lines.append("发生紧急购电的天数: %d / %d" % (n_em_days, G.shape[0]))
    return "\n".join(lines)


def table3(res):
    """表3：指定日期的紧急购电量。"""
    e = res.get('e')
    if e is None:
        return "表3：紧急购电量均为 0（确定性方案）。"
    spec = ["2025.3.20", "2025.6.21", "2025.9.23", "2025.12.21"]
    lines = ["表3  微网在指定日期的紧急购电量"]
    for ds in spec:
        idx = _find_idx(ds)
        k = idx - 31                      # 输出窗口(2.1起)的相对下标
        lines.append("  %s:" % ds)
        ranges = _merge_ranges(e[k])
        if not ranges:
            lines.append("    (无紧急购电)")
        for a, b, amt in ranges:
            lines.append("    %s-%s  %.4f kWh" % (_slot_clock(a), _slot_clock(b + 1), amt))
    return "\n".join(lines)


def main():
    global PRICE, DATES, LOAD, PV
    d1 = load_fj1()
    PRICE = d1['price']
    DATES, LOAD, PV = load_fj2()
    assert LOAD.shape[0] == 365, LOAD.shape

    det = run_deterministic()
    rob = run_robust()

    # 光伏预报误差统计(鲁棒方案依据)
    pv_f = build_pv_forecast()
    err = pv_f[31:] - PV[31:]
    out = []
    out.append("光伏预报(附件3 0:00)全年误差: 均值 %.2f kW, RMSE %.2f kW, 平均绝对 %.2f kW"
               % (err.mean(), np.sqrt((err ** 2).mean()), np.abs(err).mean()))
    out.append("")
    out.append(summary(det, "确定性方案(0:00已知实际, 紧急购电=0)"))
    out.append("")
    out.append(summary(rob, "鲁棒方案(0:00用附件3光伏预报, 实际揭晓后重调度+紧急购电)"))
    out.append("")
    out.append("---- 表3 (鲁棒方案) ----")
    out.append(table3(rob))
    out.append("")
    out.append("---- 表3 (确定性方案) ----")
    out.append(table3(det))

    write_result2(rob, "result2.xlsx")
    write_result2(det, "result2_deterministic.xlsx")

    open("_q2_summary.txt", "w", encoding="utf-8").write("\n".join(out))
    print("done")


if __name__ == "__main__":
    main()
