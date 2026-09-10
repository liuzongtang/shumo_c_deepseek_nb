# -*- coding: utf-8 -*-
"""问题3：滚动时域 + 调整购电 + 违约/溢价惩罚。

模型：
  - 每天 0:00 用附件3的0:00光伏预报(负载视为已知)制定计划购电策略 G_plan(144段)。
  - 6:00/12:00/18:00 用对应时刻的预报，对"当天剩余时段"滚动重优化，得到调整购电 G_adj。
  - 有效购电量 G_eff：  [0:00,6:00)用 G_plan, [6:00,12:00)用6:00调整, 依此类推。
  - 违约/溢价： dn=(G_plan-G_eff)^+ 违约电价50%； up=(G_eff-G_plan)^+ 超出部分1.5倍。
    总购电费 = Σ G_plan·price + Σ(1.5·up + 0.5·dn)·price + 5·Σ e·price。
  - 实际光伏揭晓后，储能重调度最小化缺口，剩余缺口按 5 倍紧急购电。

关键点：违约(下调计划)0.5倍惩罚无套利空间，最优解 dn=0；调整仅在预报恶化时上调(1.5倍)，
      用1.5倍"预购"替代5倍紧急购电，这是引入午间预报的价值所在。
"""
import numpy as np
import datetime as dt
from scipy.optimize import linprog
from data_loader import (load_fj1, load_fj2, load_fj3, DT, PMAX_E, E_MAX, E_MIN,
                         E_INIT, ETA, N_SLOT)

PRICE = None
DATES = None
LOAD = None
PV = None


def build_pv_forecast_stage():
    """附件3 -> {0,6,12,18: (365,144)} 各时刻预报(step-hold到10分钟)。
    0:00覆盖全天24h；6:00覆盖当天6-24h；12:00覆盖12-24h；18:00覆盖18-24h。"""
    fc = load_fj3()
    base = dt.datetime(2025, 1, 1)
    res = {0: np.zeros((365, N_SLOT)), 6: np.zeros((365, N_SLOT)),
           12: np.zeros((365, N_SLOT)), 18: np.zeros((365, N_SLOT))}
    for i in range(365):
        d = base + dt.timedelta(days=i)
        key = "%d-%d-%d" % (d.year, d.month, d.day)
        for tau in (0, 6, 12, 18):
            hourly = fc.get((key, "%d:00" % tau))
            if hourly is None:
                raise KeyError((key, tau))
            t0 = tau * 6
            for t in range(t0, N_SLOT):
                res[tau][i, t] = hourly[t // 6 - tau]
    return res


def stage_lp(price, load, pv_f, t0, E0, E1, G_plan_seg=None):
    """[t0,144) 段购电+储能 LP。
    G_plan_seg=None: 基础计划(0:00)，目标 min Σ price·G。
    否则: 调整阶段，目标 min Σ price·(1.5·up + 0.5·dn)，G = G_plan + up - dn。
    返回 dict: G,C,D,E(E[t0..144]), up,dn(调整阶段才有)。"""
    m = N_SLOT - t0
    pseg = price[t0:]
    Le = load[t0:] * DT
    PVe = pv_f[t0:] * DT
    has_pen = G_plan_seg is not None

    nG = nC = nD = m
    nE = m + 1
    offG, offC, offD, offE = 0, m, 2 * m, 3 * m
    if has_pen:
        offU, offN = offE + nE, offE + nE + m
        n = offN + m
    else:
        offU = offN = -1
        n = offE + nE

    c = np.zeros(n)
    if has_pen:
        c[offU:offU + m] = 1.5 * pseg
        c[offN:offN + m] = 0.5 * pseg
    else:
        c[offG:offG + m] = pseg

    Aeq, beq = [], []
    for j in range(m):                       # 储能动态
        row = np.zeros(n)
        row[offE + j + 1] = 1.0
        row[offE + j] = -1.0
        row[offC + j] = -ETA
        row[offD + j] = 1.0 / ETA
        Aeq.append(row); beq.append(0.0)
    row = np.zeros(n); row[offE + 0] = 1.0; Aeq.append(row); beq.append(E0)
    row = np.zeros(n); row[offE + m] = 1.0; Aeq.append(row); beq.append(E1)
    if has_pen:                              # G = G_plan + up - dn
        for j in range(m):
            row = np.zeros(n)
            row[offG + j] = 1.0
            row[offU + j] = -1.0
            row[offN + j] = 1.0
            Aeq.append(row); beq.append(G_plan_seg[j])

    Aub, bub = [], []                        # G + pv·Δt + D - C >= load·Δt
    for j in range(m):
        row = np.zeros(n)
        row[offG + j] = -1.0
        row[offD + j] = -1.0
        row[offC + j] = 1.0
        Aub.append(row); bub.append(PVe[j] - Le[j])

    lb = np.zeros(n); ub = np.full(n, np.inf)
    for j in range(m):
        ub[offC + j] = PMAX_E
        ub[offD + j] = PMAX_E
    for j in range(nE):
        lb[offE + j] = E_MIN; ub[offE + j] = E_MAX
    if has_pen:
        for j in range(m):
            ub[offU + j] = 1e9; ub[offN + j] = 1e9

    res = linprog(c, A_eq=np.array(Aeq), b_eq=np.array(beq),
                  A_ub=np.array(Aub), b_ub=np.array(bub),
                  bounds=list(zip(lb, ub)), method="highs")
    if not res.success:
        raise RuntimeError(res.message)
    out = {'G': res.x[offG:offG + m], 'C': res.x[offC:offC + m],
           'D': res.x[offD:offD + m], 'E': res.x[offE:offE + nE]}
    if has_pen:
        out['up'] = res.x[offU:offU + m]
        out['dn'] = res.x[offN:offN + m]
    return out


def recourse(G_eff, load, pv, E0=E_INIT, E1=E_INIT):
    """给定有效购电量 G_eff，对实际负载/光伏重调度储能，最小化紧急购电 e。"""
    Le = load * DT
    PVe = pv * DT
    nC = nD = N_SLOT
    nE = N_SLOT + 1
    ne = N_SLOT
    offC, offD, offE = 0, nC, nC + nD
    offe = offE + nE
    n = offe + ne

    c = np.zeros(n)
    c[offe:offe + ne] = PRICE

    Aeq, beq = [], []
    for t in range(N_SLOT):
        row = np.zeros(n)
        row[offE + t + 1] = 1.0
        row[offE + t] = -1.0
        row[offC + t] = -ETA
        row[offD + t] = 1.0 / ETA
        Aeq.append(row); beq.append(0.0)
    row = np.zeros(n); row[offE + 0] = 1.0; Aeq.append(row); beq.append(E0)
    row = np.zeros(n); row[offE + N_SLOT] = 1.0; Aeq.append(row); beq.append(E1)

    Aub, bub = [], []
    for t in range(N_SLOT):
        row = np.zeros(n)
        row[offD + t] = -1.0
        row[offC + t] = 1.0
        row[offe + t] = -1.0
        Aub.append(row); bub.append(G_eff[t] + PVe[t] - Le[t])

    lb = np.zeros(n); ub = np.full(n, np.inf)
    for t in range(N_SLOT):
        surplus = G_eff[t] + PVe[t] - Le[t]
        ub[offC + t] = min(PMAX_E, max(0.0, surplus))   # 充电只能充实际盈余
        ub[offD + t] = PMAX_E
    for t in range(nE):
        lb[offE + t] = E_MIN; ub[offE + t] = E_MAX

    res = linprog(c, A_eq=np.array(Aeq), b_eq=np.array(beq),
                  A_ub=np.array(Aub), b_ub=np.array(bub),
                  bounds=list(zip(lb, ub)), method="highs")
    if not res.success:
        raise RuntimeError(res.message)
    return {'C': res.x[offC:offC + nC], 'D': res.x[offD:offD + nD],
            'E': res.x[offE:offE + nE], 'e': res.x[offe:offe + ne]}


def run_day(d, pvf):
    """跑一天的四阶段滚动时域 + 执行。返回 G_plan/G_adj/C/D/E/e/up/dn。"""
    # 阶段0: 0:00 计划
    r0 = stage_lp(PRICE, LOAD[d], pvf[0][d], 0, E_INIT, E_INIT, None)
    G_plan = r0['G']                                   # 144
    # 阶段1: 6:00 调整
    r1 = stage_lp(PRICE, LOAD[d], pvf[6][d], 36, r0['E'][36], E_INIT, G_plan[36:])
    G6 = r1['G']                                       # 108 (槽36..143)
    # 阶段2: 12:00 调整
    r2 = stage_lp(PRICE, LOAD[d], pvf[12][d], 72, r1['E'][72 - 36], E_INIT, G_plan[72:])
    G12 = r2['G']                                      # 72 (槽72..143)
    # 阶段3: 18:00 调整
    r3 = stage_lp(PRICE, LOAD[d], pvf[18][d], 108, r2['E'][108 - 72], E_INIT, G_plan[108:])
    G18 = r3['G']                                      # 36 (槽108..143)
    # 有效购电量
    G_adj = np.concatenate([G_plan[0:36], G6[0:36], G12[0:36], G18[0:36]])
    # 执行
    rc = recourse(G_adj, LOAD[d], PV[d])
    up = np.maximum(0.0, G_adj - G_plan)
    dn = np.maximum(0.0, G_plan - G_adj)
    return {'G_plan': G_plan, 'G_adj': G_adj, 'C': rc['C'], 'D': rc['D'],
            'E': rc['E'], 'e': rc['e'], 'up': up, 'dn': dn}


def run(start=31, ndays=334):
    pvf = build_pv_forecast_stage()
    out = {'G_plan': [], 'G_adj': [], 'C': [], 'D': [], 'E': [], 'e': [],
           'up': [], 'dn': []}
    for d in range(start, start + ndays):
        r = run_day(d, pvf)
        for k in out:
            out[k].append(r[k])
    for k in out:
        out[k] = np.array(out[k])
    return out


def _slot_clock(t):
    m = t * 10
    return "%02d:%02d" % (m // 60, m % 60)


def _merge_ranges(e):
    out = []
    i = 0
    while i < N_SLOT:
        if e[i] <= 1e-9:
            i += 1
            continue
        j = i
        while j + 1 < N_SLOT and e[j + 1] > 1e-9:
            j += 1
        out.append((i, j, float(e[i:j + 1].sum())))
        i = j + 1
    return out


def write_result3(res, out_path):
    import openpyxl
    tpl = "附件/附件5/result3.xlsx"
    wb = openpyxl.load_workbook(tpl)
    ndays = res['G_plan'].shape[0]
    G_plan, G_adj, C, D, E = res['G_plan'], res['G_adj'], res['C'], res['D'], res['E']
    e, up, dn = res['e'], res['up'], res['dn']

    def fill_purchase(sheet_name, G, fee_col_vals):
        ws = wb[sheet_name]
        for k in range(ndays):
            row = 2 + k
            g = G[k]
            for i in range(N_SLOT):
                ws.cell(row=row, column=2 + i,
                        value=round(float(g[(i + 1) % N_SLOT]), 4))
            ws.cell(row=row, column=2 + N_SLOT, value=round(float(g.sum()), 4))
            ws.cell(row=row, column=3 + N_SLOT, value=round(float(fee_col_vals[k]), 4))

    # 计划购电量：col147 = 计划购电费
    plan_fee = np.sum(G_plan * PRICE[None, :], axis=1)
    fill_purchase("计划购电量", G_plan, plan_fee)
    # 调整购电量：col147 = 调整相关费 Σ(1.5up+0.5dn)·price
    adj_fee = np.sum((1.5 * up + 0.5 * dn) * PRICE[None, :], axis=1)
    fill_purchase("调整购电量", G_adj, adj_fee)

    # 充放电量
    ws = wb["充放电量"]
    ws.delete_rows(2, ws.max_row - 1)
    periods = [(0, 24, '0:00-4:00'), (24, 48, '4:00-8:00'), (48, 72, '8:00-12:00'),
               (72, 96, '12:00-16:00'), (96, 120, '16:00-20:00'), (120, 144, '20:00-24:00')]
    base = dt.datetime(2025, 2, 1)
    r = 2
    for k in range(ndays):
        for pi, (a, b, name) in enumerate(periods):
            if pi == 0:
                c1 = ws.cell(row=r, column=1, value=base + dt.timedelta(days=k))
                c1.number_format = 'mm-dd-yy'
            ws.cell(row=r, column=2, value=name)
            ws.cell(row=r, column=3, value=round(float(C[k][a:b].sum()), 4))
            ws.cell(row=r, column=4, value=round(float(D[k][a:b].sum()), 4))
            if pi == 0:
                c5 = ws.cell(row=r, column=5, value=dt.time(0, 0))
                c5.number_format = 'h:mm'
                ws.cell(row=r, column=6, value=round(float(E[k][0]), 4))
            elif pi == 1:
                c5 = ws.cell(row=r, column=5, value='24:00')
                c5.number_format = '@'
                ws.cell(row=r, column=6, value=round(float(E[k][-1]), 4))
            r += 1

    # 紧急购电量
    ws = wb["紧急购电量"]
    ws.delete_rows(2, ws.max_row - 1)
    r = 2
    for k in range(ndays):
        ranges = _merge_ranges(e[k])
        if not ranges:
            continue
        for ri, (a, b, amt) in enumerate(ranges):
            if ri == 0:
                c1 = ws.cell(row=r, column=1, value=base + dt.timedelta(days=k))
                c1.number_format = 'mm-dd-yy'
            ws.cell(row=r, column=2, value="%s-%s" % (_slot_clock(a), _slot_clock(b + 1)))
            ws.cell(row=r, column=3, value=round(amt, 4))
            r += 1

    wb.save(out_path)
    print("已写出", out_path)


def _find_idx(date_str):
    for i, s in enumerate(DATES):
        if str(s).replace('-', '.') == date_str or str(s) == date_str:
            return i
    target = dt.datetime.strptime(date_str, "%Y.%m.%d").date()
    for i, s in enumerate(DATES):
        if hasattr(s, 'date') and s.date() == target:
            return i
        if isinstance(s, (int, float)):
            continue
        try:
            if dt.datetime.strptime(str(s), "%Y-%m-%d").date() == target:
                return i
        except ValueError:
            pass
    raise KeyError(date_str)


def summary(res):
    G_plan, G_adj, e, up, dn = res['G_plan'], res['G_adj'], res['e'], res['up'], res['dn']
    plan_cost = float(np.sum(G_plan * PRICE[None, :]))
    plan_kwh = float(G_plan.sum())
    adj_cost = float(np.sum((1.5 * up + 0.5 * dn) * PRICE[None, :]))
    em_kwh = float(e.sum())
    em_cost = float(np.sum(e * PRICE[None, :])) * 5.0
    total = plan_cost + adj_cost + em_cost
    n_em = int((e.sum(axis=1) > 1e-9).sum())
    n_up = int((up.sum(axis=1) > 1e-9).sum())
    n_dn = int((dn.sum(axis=1) > 1e-9).sum())
    lines = ["===== 问题3：滚动时域 + 调整购电 ====="]
    lines.append("全年计划购电量: %.2f kWh" % plan_kwh)
    lines.append("全年计划购电费: %.2f 元" % plan_cost)
    lines.append("全年调整购电量 G_adj: %.2f kWh" % float(G_adj.sum()))
    lines.append("全年调整相关费(1.5up+0.5dn): %.2f 元" % adj_cost)
    lines.append("全年紧急购电量: %.2f kWh" % em_kwh)
    lines.append("全年紧急购电费(5倍): %.2f 元" % em_cost)
    lines.append("全年总购电费: %.2f 元" % total)
    lines.append("发生上调(1.5倍)的天数: %d / %d" % (n_up, G_plan.shape[0]))
    lines.append("发生下调(0.5倍违约)的天数: %d / %d" % (n_dn, G_plan.shape[0]))
    lines.append("发生紧急购电的天数: %d / %d" % (n_em, G_plan.shape[0]))
    return "\n".join(lines)


def table123(res):
    spec = ["2025.3.20", "2025.6.21", "2025.9.23", "2025.12.21"]
    lines = ["表1/表2/表3  指定日期购电策略(计划/调整/紧急)"]
    for ds in spec:
        idx = _find_idx(ds)
        k = idx - 31
        lines.append("  %s:" % ds)
        lines.append("    计划购电量全天: %.2f kWh, 计划购电费: %.2f 元"
                     % (res['G_plan'][k].sum(), np.dot(res['G_plan'][k], PRICE)))
        lines.append("    调整购电量全天: %.2f kWh, 调整相关费: %.2f 元"
                     % (res['G_adj'][k].sum(), np.dot(1.5 * res['up'][k] + 0.5 * res['dn'][k], PRICE)))
        ranges = _merge_ranges(res['e'][k])
        if not ranges:
            lines.append("    紧急购电: (无)")
        for a, b, amt in ranges:
            lines.append("    紧急购电: %s-%s  %.4f kWh" % (_slot_clock(a), _slot_clock(b + 1), amt))
    return "\n".join(lines)


def main():
    global PRICE, DATES, LOAD, PV
    d1 = load_fj1()
    PRICE = d1['price']
    DATES, LOAD, PV = load_fj2()
    assert LOAD.shape[0] == 365, LOAD.shape

    res = run()
    out = []
    out.append(summary(res))
    out.append("")
    out.append(table123(res))
    out.append("")
    # 对比问题2鲁棒(无午间调整)的总费用
    pv0 = build_pv_forecast_stage()[0]
    from solve_q2 import solve_day
    plan2 = 0.0
    em2 = 0.0
    for d in range(31, 365):
        rp = solve_day(PRICE, LOAD[d], pv0[d])
        Gp = rp['G']
        plan2 += float(np.dot(Gp, PRICE))
        rc = recourse(Gp, LOAD[d], PV[d])
        em2 += float(np.dot(rc['e'], PRICE)) * 5.0
    out.append("对比·问题2鲁棒(仅0:00预报, 无午间调整): 计划费 %.2f + 紧急费 %.2f = %.2f 元"
               % (plan2, em2, plan2 + em2))
    out.append("对比·问题3(0/6/12/18滚动调整): 计划费+调整费+紧急费 = %.2f 元"
               % (float(np.sum(res['G_plan'] * PRICE[None, :]))
                  + float(np.sum((1.5 * res['up'] + 0.5 * res['dn']) * PRICE[None, :]))
                  + float(np.sum(res['e'] * PRICE[None, :])) * 5.0))

    write_result3(res, "result3.xlsx")
    open("_q3_summary.txt", "w", encoding="utf-8").write("\n".join(out))
    print("done")


if __name__ == "__main__":
    main()
