# -*- coding: utf-8 -*-
"""问题4：波动电价下的问题2/问题3（全年连续储能，允许跨日套利）。

附件4电价随日变化(日内与逐日均波动)，储能应在低价时段/低价日充电、高价时段/高价日放电，
因此不再强制每天 24:00 回到 6000 kWh，而是全年连续优化(首末 6000 kWh 保证无"凭空能量")。

产出：
  - 确定性(完美预见，实际光伏)全局LP —— 理论下界。
  - 问题2鲁棒：全年0:00光伏预报做计划 + 实际光伏重调度最小化紧急购电(连续储能)。
  - 问题3滚动：0/6/12/18 预报滚动调整(违约/溢价) + 实际光伏重调度(连续储能)。
"""
import numpy as np
import datetime as dt
from scipy.optimize import linprog
from scipy.sparse import coo_matrix
from data_loader import (load_fj2, load_fj3, load_fj4, DT, PMAX_E, E_MAX, E_MIN,
                         E_INIT, ETA, N_SLOT)
from solve_q3 import stage_lp, build_pv_forecast_stage

PRICE_MAT = None      # 附件4 (365,144)
DATES = None
LOAD = None           # 附件2 (365,144)
PV = None             # 附件2 (365,144)


def _global_lp(price_mat, load_mat, pv_mat):
    """全年连续储能 LP(平衡用不等式允许弃光)。输入均为 (ndays,144)。
    返回 G,C,D (ndays,144), E (ndays*144+1), cost。"""
    ndays = load_mat.shape[0]
    N = ndays * N_SLOT
    Lf = load_mat.ravel() * DT
    PVf = pv_mat.ravel() * DT
    Pf = price_mat.ravel()
    offG, offC, offD, offE = 0, N, 2 * N, 3 * N
    n = 4 * N + 1

    c = np.zeros(n); c[offG:offG + N] = Pf

    rows, cols, vals = [], [], []
    b_eq = []
    for t in range(N):                     # SOC 动态
        rows += [t, t, t, t]
        cols += [offE + t + 1, offE + t, offC + t, offD + t]
        vals += [1.0, -1.0, -ETA, 1.0 / ETA]
        b_eq.append(0.0)
    rows += [N, N + 1]; cols += [offE, offE + N]; vals += [1.0, 1.0]
    b_eq += [E_INIT, E_INIT]
    A_eq = coo_matrix((vals, (rows, cols)), shape=(N + 2, n)).tocsr()

    rows2, cols2, vals2 = [], [], []
    b_ub = []
    for t in range(N):                     # G+D-C >= (L-PV)·Δt
        rows2 += [t, t, t]; cols2 += [offG + t, offD + t, offC + t]
        vals2 += [-1.0, -1.0, 1.0]
        b_ub.append(PVf[t] - Lf[t])
    A_ub = coo_matrix((vals2, (rows2, cols2)), shape=(N, n)).tocsr()

    lb = np.zeros(n); ub = np.full(n, np.inf)
    ub[offC:offC + N] = PMAX_E; ub[offD:offD + N] = PMAX_E
    lb[offE:offE + N + 1] = E_MIN; ub[offE:offE + N + 1] = E_MAX

    res = linprog(c, A_eq=A_eq, b_eq=np.array(b_eq), A_ub=A_ub, b_ub=np.array(b_ub),
                  bounds=list(zip(lb, ub)), method="highs")
    if not res.success:
        raise RuntimeError(res.message)
    G = res.x[offG:offG + N].reshape(ndays, N_SLOT)
    C = res.x[offC:offC + N].reshape(ndays, N_SLOT)
    D = res.x[offD:offD + N].reshape(ndays, N_SLOT)
    E = res.x[offE:offE + N + 1]
    return {'G': G, 'C': C, 'D': D, 'E': E, 'cost': float(np.dot(res.x[offG:offG + N], Pf))}


def _global_recourse(G_plan, price_mat, load_mat, pv_mat):
    """给定固定计划购电 G_plan(ndays,144)，对实际光伏全年连续重调度储能，最小化紧急购电。
    返回 C,D (ndays,144), E (ndays*144+1), e (ndays,144)。"""
    ndays = load_mat.shape[0]
    N = ndays * N_SLOT
    Lf = load_mat.ravel() * DT
    PVf = pv_mat.ravel() * DT
    Pf = price_mat.ravel()
    Gf = G_plan.ravel()
    offC, offD, offE = 0, N, 2 * N
    offe = 3 * N + 1                     # E 占 N+1 个变量(2N..3N)，e 从 3N+1 开始
    n = 4 * N + 1                        # C,D,E(N+1),e

    c = np.zeros(n); c[offe:offe + N] = Pf

    rows, cols, vals = [], [], []
    b_eq = []
    for t in range(N):
        rows += [t, t, t, t]
        cols += [offE + t + 1, offE + t, offC + t, offD + t]
        vals += [1.0, -1.0, -ETA, 1.0 / ETA]
        b_eq.append(0.0)
    rows += [N, N + 1]; cols += [offE, offE + N]; vals += [1.0, 1.0]
    b_eq += [E_INIT, E_INIT]
    A_eq = coo_matrix((vals, (rows, cols)), shape=(N + 2, n)).tocsr()

    rows2, cols2, vals2 = [], [], []
    b_ub = []
    for t in range(N):                   # G_plan+PV+D-C+e >= L  =>  -D+C-e <= G_plan+PV-L
        rows2 += [t, t, t]; cols2 += [offD + t, offC + t, offe + t]
        vals2 += [-1.0, 1.0, -1.0]
        b_ub.append(Gf[t] + PVf[t] - Lf[t])
    A_ub = coo_matrix((vals2, (rows2, cols2)), shape=(N, n)).tocsr()

    lb = np.zeros(n); ub = np.full(n, np.inf)
    for t in range(N):
        surplus = Gf[t] + PVf[t] - Lf[t]
        ub[offC + t] = min(PMAX_E, max(0.0, surplus))   # 充电只能充实际盈余
        ub[offD + t] = PMAX_E
    lb[offE:offE + N + 1] = E_MIN; ub[offE:offE + N + 1] = E_MAX

    res = linprog(c, A_eq=A_eq, b_eq=np.array(b_eq), A_ub=A_ub, b_ub=np.array(b_ub),
                  bounds=list(zip(lb, ub)), method="highs")
    if not res.success:
        raise RuntimeError(res.message)
    C = res.x[offC:offC + N].reshape(ndays, N_SLOT)
    D = res.x[offD:offD + N].reshape(ndays, N_SLOT)
    E = res.x[offE:offE + N + 1]
    e = res.x[offe:offe + N].reshape(ndays, N_SLOT)
    return {'C': C, 'D': D, 'E': E, 'e': e}


def run_deterministic():
    """完美预见：全年连续优化(实际光伏)。"""
    return _global_lp(PRICE_MAT[31:], LOAD[31:], PV[31:])


def run_robust(pv0):
    """问题2鲁棒：全年0:00光伏预报做计划 + 实际光伏重调度。"""
    plan = _global_lp(PRICE_MAT[31:], LOAD[31:], pv0[31:])
    rc = _global_recourse(plan['G'], PRICE_MAT[31:], LOAD[31:], PV[31:])
    plan.update({'e': rc['e'], 'C': rc['C'], 'D': rc['D'], 'E': rc['E']})
    return plan


def run_rolling(pvf, E_plan):
    """问题3滚动：0/6/12/18 预报滚动调整 + 全年连续重调度。
    E_plan: 全局预报计划的 SOC 轨迹(ndays*144+1)，提供跨日边界电量。"""
    ndays = 334
    G_plan = np.zeros((ndays, N_SLOT))
    G_adj = np.zeros((ndays, N_SLOT))
    for d in range(ndays):
        dd = 31 + d
        p = PRICE_MAT[dd]
        ld = LOAD[dd]
        E0b = E_plan[d * N_SLOT]             # 当天0:00边界电量
        E1b = E_plan[(d + 1) * N_SLOT]       # 当天24:00边界电量
        r0 = stage_lp(p, ld, pvf[0][dd], 0, E0b, E1b, None)
        G_plan[d] = r0['G']
        r1 = stage_lp(p, ld, pvf[6][dd], 36, r0['E'][36], E1b, G_plan[d][36:])
        G6 = r1['G']
        r2 = stage_lp(p, ld, pvf[12][dd], 72, r1['E'][72 - 36], E1b, G_plan[d][72:])
        G12 = r2['G']
        r3 = stage_lp(p, ld, pvf[18][dd], 108, r2['E'][108 - 72], E1b, G_plan[d][108:])
        G18 = r3['G']
        G_adj[d] = np.concatenate([G_plan[d][0:36], G6[0:36], G12[0:36], G18[0:36]])
    rc = _global_recourse(G_adj, PRICE_MAT[31:], LOAD[31:], PV[31:])
    up = np.maximum(0.0, G_adj - G_plan)
    dn = np.maximum(0.0, G_plan - G_adj)
    return {'G_plan': G_plan, 'G_adj': G_adj, 'up': up, 'dn': dn,
            'C': rc['C'], 'D': rc['D'], 'E': rc['E'], 'e': rc['e']}


def _slot_clock(t):
    m = t * 10
    return "%02d:%02d" % (m // 60, m % 60)


def _merge_ranges(e):
    out = []
    i = 0
    while i < N_SLOT:
        if e[i] <= 1e-9:
            i += 1
            continue
        j = i
        while j + 1 < N_SLOT and e[j + 1] > 1e-9:
            j += 1
        out.append((i, j, float(e[i:j + 1].sum())))
        i = j + 1
    return out


def _write_result4_2(res, out_path):
    """按 result2 模板写波动电价问题2结果。res: G,C,D,E,e (ndays,144)。"""
    import openpyxl
    wb = openpyxl.load_workbook("附件/附件5/result2.xlsx")
    ndays = res['G'].shape[0]
    G, C, D, E, e = res['G'], res['C'], res['D'], res['E'], res['e']
    ws = wb["计划购电量"]
    for k in range(ndays):
        row = 2 + k
        g = G[k]
        p = PRICE_MAT[31 + k]
        for i in range(N_SLOT):
            ws.cell(row=row, column=2 + i, value=round(float(g[(i + 1) % N_SLOT]), 4))
        ws.cell(row=row, column=2 + N_SLOT, value=round(float(g.sum()), 4))
        ws.cell(row=row, column=3 + N_SLOT, value=round(float(np.dot(g, p)), 4))
    ws = wb["充放电量"]
    ws.delete_rows(2, ws.max_row - 1)
    periods = [(0, 24, '0:00-4:00'), (24, 48, '4:00-8:00'), (48, 72, '8:00-12:00'),
               (72, 96, '12:00-16:00'), (96, 120, '16:00-20:00'), (120, 144, '20:00-24:00')]
    base = dt.datetime(2025, 2, 1)
    r = 2
    for k in range(ndays):
        Eday = E[k * N_SLOT:(k + 1) * N_SLOT + 1]
        for pi, (a, b, name) in enumerate(periods):
            if pi == 0:
                c1 = ws.cell(row=r, column=1, value=base + dt.timedelta(days=k))
                c1.number_format = 'mm-dd-yy'
            ws.cell(row=r, column=2, value=name)
            ws.cell(row=r, column=3, value=round(float(C[k][a:b].sum()), 4))
            ws.cell(row=r, column=4, value=round(float(D[k][a:b].sum()), 4))
            if pi == 0:
                c5 = ws.cell(row=r, column=5, value=dt.time(0, 0))
                c5.number_format = 'h:mm'
                ws.cell(row=r, column=6, value=round(float(Eday[0]), 4))
            elif pi == 1:
                c5 = ws.cell(row=r, column=5, value='24:00')
                c5.number_format = '@'
                ws.cell(row=r, column=6, value=round(float(Eday[-1]), 4))
            r += 1
    ws = wb["紧急购电量"]
    ws.delete_rows(2, ws.max_row - 1)
    r = 2
    for k in range(ndays):
        for ri, (a, b, amt) in enumerate(_merge_ranges(e[k])):
            if ri == 0:
                c1 = ws.cell(row=r, column=1, value=base + dt.timedelta(days=k))
                c1.number_format = 'mm-dd-yy'
            ws.cell(row=r, column=2, value="%s-%s" % (_slot_clock(a), _slot_clock(b + 1)))
            ws.cell(row=r, column=3, value=round(amt, 4))
            r += 1
    wb.save(out_path)
    print("已写出", out_path)


def _write_result4_3(res, out_path):
    """按 result3 模板写波动电价问题3结果。res: G_plan,G_adj,up,dn,C,D,E,e。"""
    import openpyxl
    wb = openpyxl.load_workbook("附件/附件5/result3.xlsx")
    ndays = res['G_plan'].shape[0]
    G_plan, G_adj, C, D, E, e = res['G_plan'], res['G_adj'], res['C'], res['D'], res['E'], res['e']
    up, dn = res['up'], res['dn']

    def fill(sheet, G, fee):
        ws = wb[sheet]
        for k in range(ndays):
            row = 2 + k
            g = G[k]
            for i in range(N_SLOT):
                ws.cell(row=row, column=2 + i, value=round(float(g[(i + 1) % N_SLOT]), 4))
            ws.cell(row=row, column=2 + N_SLOT, value=round(float(g.sum()), 4))
            ws.cell(row=row, column=3 + N_SLOT, value=round(float(fee[k]), 4))

    plan_fee = np.sum(G_plan * PRICE_MAT[31:], axis=1)
    adj_fee = np.sum((1.5 * up + 0.5 * dn) * PRICE_MAT[31:], axis=1)
    fill("计划购电量", G_plan, plan_fee)
    fill("调整购电量", G_adj, adj_fee)

    ws = wb["充放电量"]
    ws.delete_rows(2, ws.max_row - 1)
    periods = [(0, 24, '0:00-4:00'), (24, 48, '4:00-8:00'), (48, 72, '8:00-12:00'),
               (72, 96, '12:00-16:00'), (96, 120, '16:00-20:00'), (120, 144, '20:00-24:00')]
    base = dt.datetime(2025, 2, 1)
    r = 2
    for k in range(ndays):
        Eday = E[k * N_SLOT:(k + 1) * N_SLOT + 1]
        for pi, (a, b, name) in enumerate(periods):
            if pi == 0:
                c1 = ws.cell(row=r, column=1, value=base + dt.timedelta(days=k))
                c1.number_format = 'mm-dd-yy'
            ws.cell(row=r, column=2, value=name)
            ws.cell(row=r, column=3, value=round(float(C[k][a:b].sum()), 4))
            ws.cell(row=r, column=4, value=round(float(D[k][a:b].sum()), 4))
            if pi == 0:
                c5 = ws.cell(row=r, column=5, value=dt.time(0, 0))
                c5.number_format = 'h:mm'
                ws.cell(row=r, column=6, value=round(float(Eday[0]), 4))
            elif pi == 1:
                c5 = ws.cell(row=r, column=5, value='24:00')
                c5.number_format = '@'
                ws.cell(row=r, column=6, value=round(float(Eday[-1]), 4))
            r += 1
    ws = wb["紧急购电量"]
    ws.delete_rows(2, ws.max_row - 1)
    r = 2
    for k in range(ndays):
        for ri, (a, b, amt) in enumerate(_merge_ranges(e[k])):
            if ri == 0:
                c1 = ws.cell(row=r, column=1, value=base + dt.timedelta(days=k))
                c1.number_format = 'mm-dd-yy'
            ws.cell(row=r, column=2, value="%s-%s" % (_slot_clock(a), _slot_clock(b + 1)))
            ws.cell(row=r, column=3, value=round(amt, 4))
            r += 1
    wb.save(out_path)
    print("已写出", out_path)


def _cost_breakdown(res, price_mat, robust=True):
    """返回 (计划费, 紧急费(5x), 调整费, 总费) 元。"""
    price = price_mat[31:]
    e = res.get('e', None)
    if robust:
        G = res['G']
        plan = float(np.sum(G * price))
        em = float(np.sum(e * price)) * 5.0 if e is not None else 0.0
        return plan, em, 0.0, plan + em
    else:
        plan = float(np.sum(res['G_plan'] * price))
        adj = float(np.sum((1.5 * res['up'] + 0.5 * res['dn']) * price))
        em = float(np.sum(res['e'] * price)) * 5.0
        return plan, em, adj, plan + adj + em


def main():
    global PRICE_MAT, DATES, LOAD, PV
    DATES, LOAD, PV = load_fj2()
    _, PRICE_MAT = load_fj4()
    assert PRICE_MAT.shape == (365, 144), PRICE_MAT.shape
    pvf = build_pv_forecast_stage()

    det = run_deterministic()
    rob = run_robust(pvf[0])
    # 滚动所需的全局预报计划 SOC 轨迹
    plan_fc = _global_lp(PRICE_MAT[31:], LOAD[31:], pvf[0][31:])
    roll = run_rolling(pvf, plan_fc['E'])

    out = []
    out.append("===== 问题4 波动电价(全年连续储能) =====")
    out.append("附件4 电价: 全局 min=%.3f max=%.3f mean=%.4f"
               % (PRICE_MAT.min(), PRICE_MAT.max(), PRICE_MAT.mean()))
    out.append("")
    # 确定性
    p, e_, a_, t_ = _cost_breakdown(det, PRICE_MAT, robust=True)
    out.append("[确定性·完美预见] 全年购电费 = %.2f 元 (理论下界)" % t_)
    out.append("")
    # 鲁棒问题2
    p2, e2, a2, t2 = _cost_breakdown(rob, PRICE_MAT, robust=True)
    out.append("[问题2·鲁棒] 计划费 %.2f + 紧急费 %.2f = %.2f 元"
               % (p2, e2, t2))
    out.append("  紧急购电量 %.2f kWh, 紧急天数 %d/334"
               % (rob['e'].sum(), (rob['e'].sum(axis=1) > 1e-9).sum()))
    out.append("")
    # 滚动问题3
    p3, e3, a3, t3 = _cost_breakdown(roll, PRICE_MAT, robust=False)
    out.append("[问题3·滚动] 计划费 %.2f + 调整费 %.2f + 紧急费 %.2f = %.2f 元"
               % (p3, a3, e3, t3))
    out.append("  紧急购电量 %.2f kWh, 紧急天数 %d/334, 上调天数 %d"
               % (roll['e'].sum(), (roll['e'].sum(axis=1) > 1e-9).sum(),
                  (roll['up'].sum(axis=1) > 1e-9).sum()))
    out.append("")
    # 对比：附件1每日同价(连续储能，见 solve_q23_continuous.py) vs 附件4波动价
    out.append("对比(问题2鲁棒): 附件1同价总费 14725566.34 元 vs 附件4波动价总费 %.2f 元" % t2)
    out.append("对比(问题3滚动): 附件1同价总费 14562033.99 元 vs 附件4波动价总费 %.2f 元" % t3)
    out.append("")
    # 跨日储能电量利用率
    Esoc = rob['E']
    out.append("储能SOC 范围: min=%.1f max=%.1f (说明跨日搬移是否发生)"
               % (Esoc.min(), Esoc.max()))

    _write_result4_2(rob, "result4-2.xlsx")
    _write_result4_3(roll, "result4-3.xlsx")
    open("_q4_summary.txt", "w", encoding="utf-8").write("\n".join(out))
    print("done")


if __name__ == "__main__":
    main()
